"""조류 시트의 옛 형식 탭을 AviList v2026 형식으로 바꾼다(이미 'v2026'인 탭과 번역 정리 탭은 건드리지 않음).

v2026 형식(편집진이 먼저 바꾼 탭과 같음):
  RANK · ORDER · FAMILY · SCIENTIFIC NAME · ENGLISH NAME: AVILIST · 명칭 제안 · 덤해 제안 · 뿌뿌까 제안 · 검등할 제안 · 해설
  행은 family → genus → species, AviList v2026 순서. 배경색 없음.

옛 탭 → 새 탭
  - 옛 「야조회 합의본」 → 명칭 제안, 「붉부까 제안」 → 뿌뿌까 제안(같은 자리), 머리글 없는 열 → 해설
  - 옛 학명(AviList v2025 등) → v2026 학명: 같은 학명 → 두 판의 Avibase ID → 같은 속 종소명 어간
  - v2026에 새로 생긴 종(분리·신종)은 같은 속(없으면 같은 과)이 가장 많이 있는 탭에 넣고, 해설에 AviList 변경 내용을 적는다
  - v2026에서 사라진 옛 행(통합 등)은 「v2026 변경 대조」 탭에 국명·제안과 함께 모아 둔다(지우지 않음)
  - 다 바꾼 탭은 이름 끝에 ' v2026'을 붙인다

    python scripts/convert_birds_v2026.py            # 미리보기만(data/v2026_preview.json, 탭별 개수)
    python scripts/convert_birds_v2026.py --apply    # 시트에 쓴다(먼저 data/birds.xlsx 를 .secrets/backup 에 복사해 둘 것)
"""
import difflib
import json
import re
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import requests

from sheets_api import API, Sheet

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
HEADER = ["RANK", "ORDER", "FAMILY", "SCIENTIFIC NAME", "ENGLISH NAME: AVILIST", "명칭 제안", "덤해 제안",
          "뿌뿌까 제안", "검등할 제안", "해설"]
SKIP = {"남극해 조류 번역", "남극해 조류 스플릿 정리", "캐나다 조류 번역", "v2026 변경 대조"}
FAM_RE = re.compile(r"^([A-Z][a-z]+idae)\b\s*(?:\(([^)]*)\))?")
CHANGE_TAB = "v2026 변경 대조"


def stem(ep):
    for end in ("us", "um", "a", "is", "e", "i"):
        if ep.endswith(end) and len(ep) - len(end) >= 3:
            return ep[: -len(end)]
    return ep


def load_avilist(path, cols):
    rows = list(openpyxl.load_workbook(path, read_only=True).worksheets[0].iter_rows(values_only=True))
    h = [str(c).strip() if c else "" for c in rows[0]]
    ix = {k: h.index(v) for k, v in cols.items() if v in h}
    out = []
    for r in rows[1:]:
        g = lambda k: (str(r[ix[k]]).strip() if k in ix and r[ix[k]] is not None else None)
        if g("rank"):
            out.append({k: g(k) for k in cols})
    return out


def read_old_tab(values):
    """옛 탭 → (머리글, 행 목록). 행 = dict(kind, sci, cons, props[3], why, extra{})"""
    h = values[0] + [""] * (12 - len(values[0]))
    i_sci = next(i for i, c in enumerate(h) if c.startswith("학명"))
    i_cons = next(i for i, c in enumerate(h) if "합의" in c)
    i_props = [i for i, c in enumerate(h) if c.endswith("제안") and "합의" not in c][:3]
    i_why = next((i for i in range(i_cons + 1, 9) if not h[i].strip()), None)
    extra = {i: h[i] for i, c in enumerate(h) if c.strip() and i not in (i_sci, i_sci + 1, i_cons, *i_props) and i != i_why}
    rows, fam = [], None
    for r in values[1:]:
        r = r + [""] * (12 - len(r))
        a = r[i_sci].strip()
        if not a:
            continue
        if FAM_RE.match(a):
            fam = FAM_RE.match(a).group(1)
            rows.append(dict(kind="family", sci=fam))
            continue
        w = a.split()
        kind = {1: "genus", 2: "species"}.get(len(w), "subspecies")
        rows.append(dict(kind=kind, sci=" ".join(w), en=r[i_sci + 1].strip(), cons=r[i_cons].strip(), props=[r[i].strip() for i in i_props],
                         why=r[i_why].strip() if i_why is not None else "", fam=fam,
                         extra={h[i]: r[i].strip() for i in extra if r[i].strip()}))
    return rows, list(extra.values())


def main(apply):
    a25 = load_avilist(DATA / "avilist2025.xlsx", dict(rank="Taxon_rank", sci="Scientific_name", avibase="AvibaseID",
                                                         family="Family"))
    fam25 = {r["sci"]: r["family"] for r in a25 if r["rank"] == "species"}
    fam_of_genus25 = {r["sci"].split()[0]: r["family"] for r in a25 if r["rank"] == "species"}
    a26 = load_avilist(DATA / "avilist2026.xlsx", {"rank": "RANK", "order": "ORDER", "family": "FAMILY",
                                                    "sci": "SCIENTIFIC NAME", "en": "ENGLISH NAME: AVILIST",
                                                    "avibase": "AVIBASE ID", "change": "v2026 CHANGES"})
    sp26 = {r["sci"]: r for r in a26 if r["rank"] == "species"}
    # AviList v2026 변경 탭의 종 통합: v2025 학명 → v2026 학명 (종 단위 절만)
    lump25, notes = {}, defaultdict(list)
    ch = list(openpyxl.load_workbook(DATA / "avilist2026.xlsx", read_only=True)["Changes from previous version"].iter_rows(values_only=True))
    for r in ch:
        if r[0] and str(r[0]).startswith("SUBSPECIES"):
            break
        if len(r) > 5 and r[3] and r[5] and len(str(r[3]).split()) == 2 and len(str(r[5]).split()) == 2 and r[3] != r[5]:
            lump25.setdefault(str(r[5]).strip(), str(r[3]).strip())
    by_avb26 = {r["avibase"]: r["sci"] for r in a26 if r["rank"] == "species" and r["avibase"]}
    by_en26 = defaultdict(list)
    for r in a26:
        if r["rank"] == "species" and r.get("en"):
            by_en26[r["en"].lower()].append(r["sci"])
    avb25 = {r["sci"]: r["avibase"] for r in a25 if r["rank"] == "species"}
    genus26 = defaultdict(list)
    for s in sp26:
        genus26[s.split()[0]].append(s)

    def to26(sci, fam_hint=None, en=None):
        if sci in sp26:
            return sci, "같음"
        if avb25.get(sci) in by_avb26:
            return by_avb26[avb25[sci]], "Avibase ID"
        if en and len(by_en26.get(en.lower(), [])) == 1:  # 영명이 하나로 맞으면(속 이동·학명 정정)
            return by_en26[en.lower()][0], "영명"
        g, e = sci.split()
        same = [s for s in genus26.get(g, []) if stem(s.split()[1]) == stem(e)]
        if len(same) == 1:
            return same[0], "종소명 어간"
        # 학명 오타(Curvus enca, Corvus benetti 등): 비슷한 학명 하나
        near = difflib.get_close_matches(sci, list(sp26), 1, 0.9)
        if near:
            return near[0], "철자"
        # 속을 옮긴 종(Milvago chimachima → Daptrius chimachima): 같은 과의 종소명(어간)이 하나뿐일 때
        fam = fam25.get(sci) or fam_of_genus25.get(g) or fam_hint
        cands = [s for s in sp26 if stem(s.split()[1]) == stem(e) and (not fam or sp26[s]["family"] == fam)]
        if len(cands) == 1:
            return cands[0], "속 이동"
        if fam:  # 같은 과 안에서 종소명 철자가 비슷한 것 하나(Phalcoboenus australiss → Daptrius australis)
            in_fam = {s.split()[1]: s for s in sp26 if sp26[s]["family"] == fam}
            near = difflib.get_close_matches(e, list(in_fam), 1, 0.85)
            if near:
                return in_fam[near[0]], "속 이동"
        return None, None

    sh = Sheet("birds")
    meta = sh.meta()
    tabs = [(p["properties"]["title"], p["properties"]["sheetId"], p["properties"]["gridProperties"]) for p in meta["sheets"]]
    done26 = set()  # 이미 v2026 탭에 있는 종
    old = {}
    for title, sid, grid in tabs:
        if title in SKIP or title.startswith("과단위"):
            continue
        vals = sh.values(title)
        if not vals:
            continue
        if vals[0] and vals[0][0] == "RANK":
            done26 |= {r[3] for r in vals[1:] if len(r) > 3 and r[0] == "species"}
            continue
        old[title] = (sid, grid, *read_old_tab(vals))
        time.sleep(0.3)

    # 옛 종 → v2026 종, v2026 종 → 탭
    tab_of, rec_of, gone, renamed = {}, {}, [], []
    genus_rows = defaultdict(dict)  # 탭 → 속 → 옛 속 행
    for title, (_, _, rows, _) in old.items():
        for r in rows:
            if r["kind"] == "genus":
                genus_rows[title][r["sci"]] = r
            if r["kind"] != "species":
                if r["kind"] == "subspecies":
                    gone.append(dict(tab=title, **r, why_gone="아종 행"))
                continue
            new, how = to26(r["sci"], r.get("fam"), r.get("en"))
            if not new and r["sci"] in lump25:  # AviList v2026 에서 다른 종에 통합
                tgt = lump25[r["sci"]]
                gone.append(dict(tab=title, **r, why_gone=f"{tgt}에 통합(AviList v2026)"))
                notes[tgt].append(f"{r['sci']}" + (f"({r['cons']})" if r["cons"] else ""))
                continue
            if not new or new in done26:
                gone.append(dict(tab=title, **r, why_gone="v2026에 없음(통합·삭제 추정)" if not new else "이미 v2026 탭에 있음"))
                continue
            if new in rec_of:  # 두 옛 종이 한 v2026 종으로(통합)
                gone.append(dict(tab=title, **r, why_gone=f"{new}에 통합"))
                continue
            tab_of[new], rec_of[new] = title, r
            if new != r["sci"]:
                renamed.append(dict(tab=title, old=r["sci"], new=new, how=how, cons=r["cons"]))
    # 새 종(분리·신종): 같은 속 → 같은 과가 가장 많은 탭
    fam26 = {s: r["family"] for s, r in sp26.items()}
    new_sp = [s for s in sp26 if s not in tab_of and s not in done26]
    placed_new = []
    for s in new_sp:
        votes = Counter(tab_of[x] for x in genus26[s.split()[0]] if x in tab_of) or \
            Counter(tab_of[x] for x in sp26 if x in tab_of and fam26[x] == fam26[s])
        if votes:
            tab_of[s] = votes.most_common(1)[0][0]
            placed_new.append(s)

    # 탭마다 새 행(AviList v2026 순서)
    fam_rows = {r["sci"]: r for r in a26 if r["rank"] == "family"}
    new_tabs = defaultdict(list)
    cur_fam, cur_gen = {}, {}
    for r in a26:
        if r["rank"] != "species" or r["sci"] not in tab_of:
            continue
        t = tab_of[r["sci"]]
        g = r["sci"].split()[0]
        if cur_fam.get(t) != r["family"]:
            f = fam_rows.get(r["family"], {})
            new_tabs[t].append(["family", r["order"], r["family"], r["family"], f.get("en") or "", "", "", "", "", ""])
            cur_fam[t], cur_gen[t] = r["family"], None
        if cur_gen.get(t) != g:
            old_g = next((genus_rows[x].get(g) for x in genus_rows if genus_rows[x].get(g)), None) or {}
            p = (old_g.get("props") or []) + ["", "", ""]
            new_tabs[t].append(["genus", r["order"], r["family"], g, "", old_g.get("cons", ""), p[0], p[1], p[2],
                                old_g.get("why", "")])
            cur_gen[t] = g
        o = rec_of.get(r["sci"], {})
        p = (o.get("props") or []) + ["", "", ""]
        why = o.get("why", "")
        if notes.get(r["sci"]):
            why = (why + " " if why else "") + "[AviList v2026] 통합: " + ", ".join(notes[r["sci"]])
        if not o and r.get("change"):
            why = f"[AviList v2026] {r['change']}"
        elif o and o["sci"] != r["sci"]:
            why = (why + " " if why else "") + f"[AviList v2026] 이전 학명 {o['sci']}"
        row = ["species", r["order"], r["family"], r["sci"], r.get("en") or "", o.get("cons", ""), p[0], p[1], p[2], why]
        new_tabs[t].append(row + [o.get("extra", {}).get(k, "") for k in old[t][3]])

    report = dict(tabs={t: dict(rows=len(v), species=sum(1 for x in v if x[0] == "species")) for t, v in new_tabs.items()},
                  renamed=renamed, gone=[{k: v for k, v in g.items() if k != "extra"} for g in gone], new=placed_new,
                  unplaced_new=[s for s in new_sp if s not in tab_of])
    (DATA / "v2026_preview.json").write_text(json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"옛 형식 탭 {len(old)}개 → 새 행 {sum(len(v) for v in new_tabs.values())}줄")
    print(f"  학명 바뀐 종 {len(renamed)} · 새로 넣는 종 {len(placed_new)} · 넣을 탭 못 찾은 새 종 {len(report['unplaced_new'])}"
          f" · v2026에 없는 옛 행 {len(gone)}")
    print("  옛 탭 가운데 종이 하나도 안 남는 탭:", [t for t in old if t not in new_tabs])
    if not apply:
        return

    # 쓰기: 탭마다 내용·서식을 지우고 새 행을 쓴 뒤 이름에 ' v2026'
    for title, (sid, grid, _, extra) in old.items():
        rows = [HEADER + extra] + new_tabs.get(title, [])
        ncol = max(len(HEADER) + len(extra), grid.get("columnCount", 10))
        nrow = max(len(rows) + 20, grid.get("rowCount", 100))
        reqs = [
            {"updateSheetProperties": {"properties": {"sheetId": sid, "gridProperties": {"rowCount": nrow, "columnCount": ncol}},
                                       "fields": "gridProperties.rowCount,gridProperties.columnCount"}},
            {"unmergeCells": {"range": {"sheetId": sid}}},
            {"updateCells": {"range": {"sheetId": sid}, "fields": "userEnteredValue,userEnteredFormat"}},
            {"updateCells": {"start": {"sheetId": sid, "rowIndex": 0, "columnIndex": 0},
                             "rows": [{"values": [{"userEnteredValue": {"stringValue": str(c)}} for c in r]} for r in rows],
                             "fields": "userEnteredValue"}},
            {"repeatCell": {"range": {"sheetId": sid, "startRowIndex": 0, "endRowIndex": 1},
                            "cell": {"userEnteredFormat": {"textFormat": {"bold": True}}}, "fields": "userEnteredFormat.textFormat.bold"}},
            {"updateSheetProperties": {"properties": {"sheetId": sid, "title": f"{title} v2026",
                                                      "gridProperties": {"frozenRowCount": 1}},
                                       "fields": "title,gridProperties.frozenRowCount"}},
        ]
        requests.post(API + sh.id + ":batchUpdate", headers=sh._h(), json={"requests": reqs}, timeout=300).raise_for_status()
        print("  바꿈:", title, "→", f"{title} v2026", len(rows) - 1, "줄", flush=True)
        time.sleep(1.5)

    # v2026 변경 대조 탭(통합·삭제 등으로 빠진 옛 행, 학명이 바뀐 종)
    head = ["구분", "옛 탭", "옛 학명", "v2026 학명", "명칭 제안(옛 합의본)", "덤해 제안", "뿌뿌까 제안", "검등할 제안", "해설", "비고"]
    rows = [head]
    for g in gone:
        p = (g.get("props") or []) + ["", "", ""]
        rows.append(["v2026에 없음", g["tab"], g["sci"], "", g.get("cons", ""), p[0], p[1], p[2], g.get("why", ""), g["why_gone"]])
    for r in renamed:
        rows.append(["학명 변경", r["tab"], r["old"], r["new"], r.get("cons", ""), "", "", "", "", r["how"]])
    for s in report["unplaced_new"]:
        rows.append(["새 종(탭 미정)", "", "", s, "", "", "", "", sp26[s].get("change") or "", sp26[s]["family"]])
    if CHANGE_TAB not in [t for t, _, _ in tabs]:
        requests.post(API + sh.id + ":batchUpdate", headers=sh._h(), timeout=120,
                      json={"requests": [{"addSheet": {"properties": {"title": CHANGE_TAB,
                                                                      "gridProperties": {"rowCount": len(rows) + 10, "columnCount": 10}}}}]}).raise_for_status()
    sh.update([{"range": f"'{CHANGE_TAB}'!A1", "values": rows}])
    print("  변경 대조 탭:", len(rows) - 1, "줄")


if __name__ == "__main__":
    main("--apply" in sys.argv)
