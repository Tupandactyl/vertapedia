"""Reptile Database 종 목록을 파충류 학명의 기준으로 삼는다.

- data/rdb_live.json      : 웹사이트의 현재 전체 종 목록 → 어떤 학명이 유효한 종인지(기준)
- data/rdb_checklist.xlsx : 내려받는 checklist → 과·목·아종 (웹에서 속을 옮긴 종은 종소명으로 이어 받는다)
- data/rdb_synonyms.xlsx  : 이명 → 현재 학명 (Reptile Database가 따로 내는 이명 목록)
(모두 scripts/fetch_rdb.py 로 받는다)

시트 종을 목록에 맞춘다.
  1. 학명이 같으면 그대로
  2. 목록의 다른 종 아래 아종으로 들어가 있으면 → 아종 단위 통합 (국명은 그 아종에 붙인다)
  3. 같은 속에 종소명 어간이 같은 새 학명이 있으면 → 학명 변경(성 일치 등)
  4. 이명 목록에서 현재 학명을 찾으면 → 학명 변경, 이미 시트에 있는 종이면 통합
  5. 같은 속에 철자가 거의 같은 새 학명 → 학명 변경(철자)
  6. 다른 속에 종소명이 같은 새 학명이 하나뿐이면(여럿이면 같은 과) → 학명 변경(속 이동)
  7. 그래도 없으면 목록에서 빠진 종(검토 대상)
목록에만 있는 종은 시트 순서에서 같은 속(없으면 같은 과) 마지막 종 뒤에 넣는다.
"""
import difflib
import json
import re
from collections import Counter, defaultdict

import openpyxl

ORDER = {"Sauria": ("Squamata", None), "Serpentes": ("Squamata", "Serpentes")}  # 도마뱀아목은 두지 않음(측계통군)


def stem(ep):
    for end in ("orum", "arum", "ensis", "ense", "us", "um", "a", "is", "e", "i"):
        if ep.endswith(end) and len(ep) - len(end) >= 3:
            return ep[: -len(end)]
    return ep


def same_author(a, b):
    first = lambda x: (re.findall(r"[A-ZÀ-Þ][A-Za-zÀ-ÿ'\-]+", (x or "").upper()) or [""])[0]
    return bool(a and b) and first(a) == first(b)


def parse_checklist(path):
    rows = list(openpyxl.load_workbook(path, read_only=True)["Checklist"].iter_rows(values_only=True))
    h = [str(c).strip() if c else "" for c in rows[0]]
    ix = {k: h.index(k) for k in ("Species", "Author", "Subspecies", "order", "Family", "changes")}
    out = {}
    for r in rows[1:]:
        sci = (r[ix["Species"]] or "").strip()
        if not sci:
            continue
        ssp = []
        for line in (r[ix["Subspecies"]] or "").splitlines():
            w = line.split()
            if len(w) >= 3:
                ssp.append(dict(sci=" ".join(w[:3]), author=" ".join(w[3:]) or None))
        out[sci] = dict(sci=sci, author=r[ix["Author"]], order=r[ix["order"]], family=r[ix["Family"]],
                        changes=r[ix["changes"]], subspecies=ssp)
    return out


def current_list(data_dir):
    """웹 목록의 종마다 checklist의 과·목·아종을 붙인다. 웹에서 속을 옮긴 종은 같은 종소명의 옛 학명에서 이어 받는다."""
    ck = parse_checklist(data_dir / "rdb_checklist.xlsx")
    raw = json.loads((data_dir / "rdb_live.json").read_text(encoding="utf-8"))
    live, higher = raw["species"], raw.get("higher", {})
    gone = {s: v for s, v in ck.items() if s not in live}
    fresh = [s for s in live if s not in ck]
    by_ep = defaultdict(list)
    for s in gone:
        by_ep[stem(s.split()[1])].append(s)
    out, moved = {}, {}
    for s in live:
        if s in ck:
            out[s] = dict(ck[s], author=live[s])
    genus_info = {}
    for v in out.values():
        genus_info.setdefault(v["sci"].split()[0], (v["order"], v["family"]))
    # 같은 종소명(어간)이고, 옮겨 간 속이 이미 있는 속이면 과가 같아야 한다
    cands = {f: [g for g in by_ep.get(stem(f.split()[1]), [])
                 if f.split()[0] not in genus_info or genus_info[f.split()[0]][1] == gone[g]["family"]] for f in fresh}
    votes = Counter((g.split()[0], f.split()[0]) for f, cs in cands.items() for g in cs)
    for f in fresh:
        cs = sorted(cands[f], key=lambda g: -votes[(g.split()[0], f.split()[0])])
        g = cs[0] if cs and (len(cs) == 1 or votes[(cs[0].split()[0], f.split()[0])] > votes[(cs[1].split()[0], f.split()[0])]) else None
        if g and g not in moved:
            old = gone[g]
            moved[g] = f
            out[f] = dict(old, sci=f, author=live[f], changes=f"웹에서 학명 변경 (checklist: {g})",
                          subspecies=[dict(x, sci=" ".join([f.split()[0], f.split()[1]] + x["sci"].split()[2:]))
                                      for x in old["subspecies"]])
        else:
            order, fam = genus_info.get(f.split()[0]) or (higher.get(f.split()[0], {}).get("order"), higher.get(f.split()[0], {}).get("family"))
            out[f] = dict(sci=f, author=live[f], order=order, family=fam, changes="웹에 새로 오름", subspecies=[])
    return out, moved, [g for g in gone if g not in moved]


def parse_synonyms(path):
    syn = defaultdict(set)
    for a, b in list(openpyxl.load_workbook(path, read_only=True).worksheets[0].iter_rows(values_only=True))[1:]:
        if a and b:
            syn[a.strip()].add(b.strip())
    return syn


def apply(sheet, data_dir, label):
    """sheet(파충류+조류 분류군 목록)을 고친다. 바뀐 내역을 돌려준다."""
    rdb, moved, lumped = current_list(data_dir)
    syn = parse_synonyms(data_dir / "rdb_synonyms.xlsx")
    for g, f in moved.items():  # checklist 이후 웹에서 바뀐 학명도 이명으로 따라가게
        syn[g].add(f)
    sp = [t for t in sheet if t["group"] == "REPTILIA" and t["rank"] == "species"]
    have = {t["sci"] for t in sp}
    free = {s for s in rdb if s not in have}  # 목록에만 있는 학명(아직 시트 종이 차지하지 않음)
    ssp_of = {" ".join(s["sci"].split()[::2]): (v["sci"], s["sci"]) for v in rdb.values() for s in v["subspecies"]}
    by_key = {t["sci"]: t for t in sp}
    log = dict(renamed=[], merged=[], dropped=[], added=[], web_lumped=lumped)

    def rename(t, new, how):
        log["renamed"].append(dict(old=t["sci"], new=new, how=how, name=t.get("consensus")))
        t.setdefault("synonyms", []).append(t["sci"])
        t["review"] = f"학명 변경: 시트 {t['sci']} → {label} {new}"
        by_key.pop(t["sci"], None)
        t["sci"] = new
        by_key[new] = t
        free.discard(new)

    merged = []
    for t in sp:
        s = t["sci"]
        if s in rdb:
            continue
        g, e = s.split()[:2]
        if (g + " " + e) in ssp_of and ssp_of[g + " " + e][0] != s:
            parent, ssp_sci = ssp_of[g + " " + e]
            merged.append((t, parent, ssp_sci, "아종 단위 통합"))
            continue
        same_genus = [x for x in free if x.split()[0] == g and stem(x.split()[1]) == stem(e)]
        if len(same_genus) == 1:
            rename(t, same_genus[0], "종소명 어미")
            continue
        cur = [x for x in syn.get(s, ()) if x in rdb]
        if len(cur) == 1:
            if cur[0] in free:
                rename(t, cur[0], "이명 목록")
            else:
                merged.append((t, cur[0], None, "종 통합(이명)"))
            continue
        near = difflib.get_close_matches(e, [x.split()[1] for x in free if x.split()[0] == g], 1, 0.85)
        if near:
            rename(t, f"{g} {near[0]}", "철자")
            continue
        # 속을 옮긴 것으로 보려면 명명자(첫 성)도 같아야 한다 — 종소명만 같은 남의 종으로 가지 않게
        other = [x for x in free if x.split()[1] == e and same_author(t.get("author"), rdb[x]["author"])]
        if len(other) > 1:  # 같은 종소명이 여럿이면 같은 과 쪽
            other = [x for x in other if rdb[x]["family"] == t.get("family")]
        if len(other) == 1:
            rename(t, other[0], "속 이동")
            continue
        log["dropped"].append(dict(sci=s, family=t.get("family"), name=t.get("consensus")))
        t["dropped"] = True

    # 아종 목록·명명자는 목록을 따른다
    for t in sp:
        r = rdb.get(t["sci"])
        if r:
            t["author"] = r["author"] or t.get("author")
            t["subspecies"] = [dict(x) for x in r["subspecies"]]
            if r["changes"]:
                t["rdb_change"] = r["changes"]
            if not t.get("family") and (r["family"] or "").endswith("idae"):  # 분류미정 탭의 종: 목록의 과로
                t["family"] = r["family"]
                log.setdefault("placed", []).append(dict(sci=t["sci"], family=r["family"]))
    for t, parent, ssp_sci, how in merged:
        p = by_key[parent]
        log["merged"].append(dict(old=t["sci"], into=parent, ssp=ssp_sci, how=how, name=t.get("consensus")))
        p.setdefault("synonyms", []).append(t["sci"])
        entry = next((x for x in p.get("subspecies", []) if x["sci"] == ssp_sci), None) if ssp_sci else None
        if entry is None and t.get("consensus"):
            entry = dict(sci=ssp_sci or t["sci"])
            p.setdefault("subspecies", []).append(entry)
        if entry is not None:
            entry["note"] = f"{how} · 시트 {t['sci']}"
            if t.get("consensus"):
                entry["ko"] = t["consensus"]
        sheet.remove(t)
    for t in [t for t in sp if t.get("dropped")]:
        sheet.remove(t)

    # 목록에만 있는 종 넣기
    fam_rows = {t["sci"]: t for t in sheet if t["group"] == "REPTILIA" and t["rank"] == "family"}
    for p in log.get("placed", []):
        if p["family"] not in fam_rows:
            order, sub = ORDER.get(rdb[p["sci"]]["order"], (rdb[p["sci"]]["order"], None))
            fam_rows[p["family"]] = dict(group="REPTILIA", rank="family", sci=p["family"], en=None, consensus=None,
                                         order=order, suborder=sub, proposals={}, rationale=None, sheet=label)
            sheet.append(fam_rows[p["family"]])
    genus_fam = {}
    for t in sheet:
        if t["group"] == "REPTILIA" and t["rank"] == "species":
            genus_fam.setdefault(t["sci"].split()[0], t.get("family"))
    for s in sorted(free):
        r = rdb[s]
        g = s.split()[0]
        order, sub = ORDER.get(r["order"], (r["order"], None))
        fam = genus_fam[g] if g in genus_fam else (r["family"] if r["family"] in fam_rows or (r["family"] or "").endswith("idae") else None)
        if fam and fam not in fam_rows:
            fam_rows[fam] = dict(group="REPTILIA", rank="family", sci=fam, en=None, consensus=None, order=order,
                                 suborder=sub, proposals={}, rationale=None, sheet=label)
            sheet.append(fam_rows[fam])
        t = dict(group="REPTILIA", rank="species", sci=s, family=fam, author=r["author"],
                 order_hint=None if fam else order, suborder_hint=None if fam else sub,
                 en=None, consensus=None, proposals={}, rationale=None, sheet=label,
                 subspecies=[dict(x) for x in r["subspecies"]], rdb_change=r["changes"] or "시트에 없음")
        last_g = last_f = None
        for i, x in enumerate(sheet):
            if x["group"] == "REPTILIA" and x["rank"] == "species":
                if x["sci"].split()[0] == g:
                    last_g = i
                if fam and x.get("family") == fam:
                    last_f = i
        idx = last_g if last_g is not None else last_f
        sheet.insert(idx + 1 if idx is not None else len(sheet), t)
        genus_fam.setdefault(g, fam)
        log["added"].append(dict(sci=s, family=fam, change=r["changes"]))
    return log
